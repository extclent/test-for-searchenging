# -*- coding: utf-8 -*-
"""
比赛管理与筛选系统 - 后端服务
Flask + SQLite + openpyxl
提供 REST API 和 Excel 解析，同时静态托管前端页面
"""
import os
import sys
import json
import re
import sqlite3
from datetime import datetime, timedelta
from io import BytesIO

# 把 pip 安装到 --target 的目录加入 sys.path
PYLIBS = os.path.join(os.environ.get('TEMP', ''), 'pylibs')
if PYLIBS not in sys.path:
    sys.path.insert(0, PYLIBS)

from flask import Flask, request, jsonify, send_file, send_from_directory

# ---------------- 路径配置 ----------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
HTML_FILE = os.path.join(BASE_DIR, 'Untitled-0.html')
DB_PATH = os.path.join(BASE_DIR, 'competitions.db')

app = Flask(__name__, static_folder=None)

# ---------------- 字段映射规则（从前端迁移） ----------------
FIELD_KEYWORDS = {
    'name':       ['赛事名称', '比赛名称', '名称', '赛事', 'event', 'name', 'title'],
    'level':      ['赛事级别', '级别', '层级', 'level'],
    'signupTime': ['举办时间', '报名时间', '报名日期', '举办', '报名', '时间', '截止', 'deadline', 'date', 'signup'],
    'compTimeRaw':['比赛日期', '比赛时间', '竞赛时间', '竞赛日期', '赛事时间', '赛事日期'],
    'category':   ['参赛赛道', '赛道', '类别', '分类', 'category', 'track'],
    'organizer':  ['主办单位', '主办方', '主办', 'organizer', 'host'],
    'undertaker': ['承办单位', '承办', '执行', '组织', 'undertaker'],
    'signupLink': ['报名链接', '链接', 'link', 'url', 'website'],
}

TAG_RULES = {
    '数学建模': ['数学', '建模', 'matlab'],
    '程序设计': ['编程', '程序', '代码', '算法', 'acm', 'leetcode', '开发'],
    '人工智能': ['ai', '人工智能', '机器学习', '深度学习', '数据'],
    '机械创新': ['机械', '制造', '结构', 'cad'],
    '电子设计': ['电子', '电路', '单片机', '嵌入式'],
    '商业计划': ['商业', '创业', 'bp', '路演'],
    '数学': ['数学', '高数', '微积分'],
    '物理': ['物理', '力学', '光电'],
    '化学': ['化学', '材料', '化工'],
    '文学': ['文学', '写作', '作文', '诗歌'],
    '外语': ['英语', '日语', '翻译', '口语', '演讲'],
    '设计': ['设计', '美术', '创意', '海报', 'ui'],
    '体育': ['体育', '运动', '球', '田径'],
    '音乐': ['音乐', '歌唱', '乐器', '合唱'],
    '国家级': ['全国', '国家'],
    '省级': ['省', '全省'],
    '国际级': ['国际', '世界'],
    '高奖金': ['奖金', '万元', '万'],
    '高含金量': ['教育部', '官方', '认证', '重点'],
}

# 多赛道/多单位分隔符
SPLIT_RE = re.compile(r'[\/／、,，;；\s]+')


def normalize_header(s):
    """规范化表头：去空白、转小写"""
    return str(s or '').strip().lower().replace(' ', '')


def build_column_map(headers):
    """根据表头构建字段→列索引映射"""
    col_map = {}
    for idx, h in enumerate(headers):
        nh = normalize_header(h)
        for field, kws in FIELD_KEYWORDS.items():
            if field in col_map:
                continue
            if any(k == nh or k in nh for k in kws):
                col_map[field] = idx
    return col_map


def parse_signup_time_val(val):
    """解析报名时间值，返回 (precision, value)"""
    if not val:
        return 'date', ''
    val = str(val).strip()
    # 纯数字 → Excel 日期序列号
    if re.match(r'^\d+$', val):
        ser = int(val)
        base = datetime(1899, 12, 30)
        d = base + timedelta(days=ser)
        return 'date', f'{d.year}-{d.month:02d}-{d.day:02d}'
    # YYYY-M 或 YYYY/M
    if re.match(r'^\d{4}[-/]\d{1,2}$', val):
        return 'month', val.replace('/', '-')
    # YYYY-M-D 或 YYYY/M/D
    if re.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}$', val):
        return 'date', val.replace('/', '-')
    return 'date', val


def cell_to_str(c):
    if c is None:
        return ''
    if isinstance(c, datetime):
        return c.strftime('%Y-%m-%d')
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c)


# 级别类标签只依据赛事级别列判定，避免时间文本里的省赛/国赛误伤
LEVEL_TAGS = {'国家级', '省级', '国际级'}


def split_category(cat):
    # 拆分赛道文本：括号内的斜杠不再把外层名称切残
    out = []
    for seg in re.split(r'[、,，;；]+', str(cat)):
        seg = seg.strip()
        if not seg:
            continue
        m = re.match(r'^(.*?)[（(](.*?)[)）]\s*$', seg)
        if m:
            outer = m.group(1).strip()
            if outer:
                out.append(outer)
            for inner in re.split(r'[\/／]+', m.group(2)):
                inner = inner.strip()
                if inner:
                    out.append(inner)
        else:
            for sub in re.split(r'[\/／]+', seg):
                sub = sub.strip()
                if sub:
                    out.append(sub)
    return out


def auto_suggest_tags_for_row(row):
    # 不含报名时间文本，防止省赛/国赛误判级别
    tags = set()
    base_text = ' '.join(str(row.get(k, '')) for k in
                         ('name', 'level', 'category', 'organizer', 'undertaker')).lower()
    level_text = str(row.get('level', '')).lower()
    for tag, kws in TAG_RULES.items():
        scope = level_text if tag in LEVEL_TAGS else base_text
        if any(k.lower() in scope for k in kws):
            tags.add(tag)
    if row.get('level'):
        tags.add(str(row['level']))
    for t in split_category(row.get('category', '')):
        tags.add(t)
    return list(tags)


# ---------------- 数据库 ----------------
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    conn.execute('''
        CREATE TABLE IF NOT EXISTS competitions (
            id REAL PRIMARY KEY,
            name TEXT, level TEXT, category TEXT,
            organizer TEXT, undertaker TEXT, signup_link TEXT,
            signup_precision TEXT, signup_time TEXT, signup_time_label TEXT,
            tags TEXT,
            comp_times TEXT,
            created_at TEXT
        )
    ''')
    # 兼容已有库：若旧表缺少 comp_times 列则自动添加
    cols = {r['name'] for r in conn.execute('PRAGMA table_info(competitions)').fetchall()}
    if 'comp_times' not in cols:
        conn.execute('ALTER TABLE competitions ADD COLUMN comp_times TEXT')
    conn.commit()
    conn.close()


def normalize_name(name):
    # 业务唯一键：移除全部空白（含全角空格）、转小写
    t = re.sub(r'\s+', '', str(name or '').replace('\u3000', '')).lower()
    return t


def existing_name_set(conn):
    rows = conn.execute('SELECT name FROM competitions').fetchall()
    return {normalize_name(r['name']) for r in rows if r['name']}


def row_to_dict(row):
    """把数据库行转为前端 competition 对象"""
    try:
        comp_times = json.loads(row['comp_times']) if row['comp_times'] else {}
    except (ValueError, TypeError):
        comp_times = {}
    return {
        'id': row['id'],
        'name': row['name'] or '',
        'level': row['level'] or '',
        'category': row['category'] or '',
        'organizer': row['organizer'] or '',
        'undertaker': row['undertaker'] or '',
        'signupLink': row['signup_link'] or '',
        'signupPrecision': row['signup_precision'] or 'date',
        'signupTime': row['signup_time'] or '',
        'signupTimeLabel': row['signup_time_label'] or '',
        'tags': json.loads(row['tags']) if row['tags'] else [],
        'compTimes': comp_times if isinstance(comp_times, dict) else {},
        'createdAt': row['created_at'] or '',
    }


# ---------------- 路由 ----------------
@app.route('/')
def index():
    return send_file(HTML_FILE)


@app.route('/api/competitions', methods=['GET'])
def list_competitions():
    conn = get_db()
    rows = conn.execute('SELECT * FROM competitions ORDER BY id DESC').fetchall()
    conn.close()
    return jsonify([row_to_dict(r) for r in rows])


@app.route('/api/competitions', methods=['POST'])
def create_competition():
    data = request.get_json(force=True)
    comp_times = data.get('compTimes', {})
    if not isinstance(comp_times, dict):
        comp_times = {}
    comp = {
        'id': data.get('id') or (datetime.now().timestamp() * 1000),
        'name': data.get('name', ''),
        'level': data.get('level', ''),
        'category': data.get('category', ''),
        'organizer': data.get('organizer', ''),
        'undertaker': data.get('undertaker', ''),
        'signup_link': data.get('signupLink', ''),
        'signup_precision': data.get('signupPrecision', 'date'),
        'signup_time': data.get('signupTime', ''),
        'signup_time_label': data.get('signupTimeLabel', ''),
        'tags': json.dumps(data.get('tags', []), ensure_ascii=False),
        'comp_times': json.dumps(comp_times, ensure_ascii=False),
        'created_at': data.get('createdAt', datetime.now().strftime('%Y/%m/%d %p%I:%M:%S')),
    }
    key = normalize_name(comp['name'])
    conn = get_db()
    if key and key in existing_name_set(conn):
        conn.close()
        return jsonify({'ok': False, 'duplicate': True,
                        'message': '该比赛已存在，未重复录入',
                        'existingName': comp['name']}), 409
    conn.execute('''
        INSERT INTO competitions
        (id, name, level, category, organizer, undertaker, signup_link,
         signup_precision, signup_time, signup_time_label, tags, comp_times, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (comp['id'], comp['name'], comp['level'], comp['category'],
          comp['organizer'], comp['undertaker'], comp['signup_link'],
          comp['signup_precision'], comp['signup_time'], comp['signup_time_label'],
          comp['tags'], comp['comp_times'], comp['created_at']))
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'id': comp['id']})


@app.route('/api/competitions/<float:sid>', methods=['DELETE'])
@app.route('/api/competitions/<int:sid>', methods=['DELETE'])
def delete_competition(sid):
    conn = get_db()
    cur = conn.execute('DELETE FROM competitions WHERE id = ?', (sid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'deleted': cur.rowcount})


@app.route('/api/competitions/batch', methods=['POST'])
def batch_create():
    data = request.get_json(force=True)
    items = data.get('items', [])
    conn = get_db()
    seen = existing_name_set(conn)  # 库内已有
    added = 0
    skipped = []
    for item in items:
        key = normalize_name(item.get('name', ''))
        if not key:
            continue
        if key in seen:
            skipped.append(item.get('name', ''))
            continue
        seen.add(key)
        cid = item.get('id') or (datetime.now().timestamp() * 1000 + added * 0.001)
        ct = item.get('compTimes', {})
        if not isinstance(ct, dict):
            ct = {}
        conn.execute('''
            INSERT INTO competitions
            (id, name, level, category, organizer, undertaker, signup_link,
             signup_precision, signup_time, signup_time_label, tags, comp_times, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (cid, item.get('name', ''), item.get('level', ''),
              item.get('category', ''), item.get('organizer', ''),
              item.get('undertaker', ''), item.get('signupLink', ''),
              item.get('signupPrecision', 'date'), item.get('signupTime', ''),
              item.get('signupTimeLabel', ''),
              json.dumps(item.get('tags', []), ensure_ascii=False),
              json.dumps(ct, ensure_ascii=False),
              item.get('createdAt', datetime.now().strftime('%Y/%m/%d %p%I:%M:%S'))))
        added += 1
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'added': added, 'skipped': len(skipped),
                    'skippedNames': skipped})


@app.route('/api/excel/parse', methods=['POST'])
def excel_parse():
    """接收 Excel 文件，解析、字段映射、推荐标签，返回预览 JSON"""
    if 'file' not in request.files:
        return jsonify({'error': '未收到文件'}), 400
    f = request.files['file']
    filename = f.filename or ''
    ext = os.path.splitext(filename)[1].lower()

    try:
        if ext == '.csv':
            # 用 utf-8-sig 兼容带 BOM 的 CSV
            content = f.read()
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('gbk', errors='replace')
            import csv as csv_mod
            all_rows = list(csv_mod.reader(text.splitlines()))
        elif ext == '.xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(f.read()), data_only=True)
            ws = wb.active
            all_rows = [[cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        elif ext == '.xls':
            import xlrd
            book = xlrd.open_workbook(file_contents=f.read())
            sh = book.sheet_by_index(0)
            all_rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            dt = xlrd.xldate_as_datetime(cell.value, book.datemode)
                            row.append(dt.strftime('%Y-%m-%d'))
                        except Exception:
                            row.append(str(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        v = cell.value
                        row.append(str(int(v)) if float(v).is_integer() else str(v))
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        row.append('')
                    else:
                        row.append(str(cell.value))
                all_rows.append(row)
        else:
            return jsonify({'error': f'不支持的文件格式 {ext}，请上传 .xlsx / .xls / .csv'}), 400

        if not all_rows:
            return jsonify({'error': '文件中未检测到数据'}), 400

        # 自动定位表头行：找到包含 name 字段且至少匹配 2 个字段的行（前 5 行内）
        header_idx = 0
        headers = []
        col_map = {}
        for h in range(min(len(all_rows), 5)):
            candidate = [str(x or '') for x in (all_rows[h] if h < len(all_rows) else [])]
            candidate_map = build_column_map(candidate)
            matched = len(candidate_map)
            if 'name' in candidate_map and matched >= 2:
                header_idx = h
                headers = candidate
                col_map = candidate_map
                break

        if not headers:
            headers = [str(x or '') for x in (all_rows[0] if all_rows else [])]
            col_map = build_column_map(headers)

        data_rows = all_rows[header_idx + 1:]

        def get_field(arr, field):
            if field in col_map and col_map[field] < len(arr):
                return arr[col_map[field]] or ''
            return ''

        conn = get_db()
        db_names = existing_name_set(conn)
        conn.close()
        file_seen = set()
        parsed = []
        for arr in data_rows:
            if not arr:
                continue
            time_raw = get_field(arr, 'signupTime')
            # 若同时存在"比赛日期"列，合并到 signupTime 末尾（保留初赛/省赛/国赛时间）
            if 'compTimeRaw' in col_map:
                comp_time_raw = get_field(arr, 'compTimeRaw')
                if comp_time_raw:
                    time_raw = (str(time_raw).rstrip() + '\n' + str(comp_time_raw).rstrip()).strip()
            precision, signup_val = parse_signup_time_val(time_raw)
            is_multiline = '\n' in str(time_raw) or '\r' in str(time_raw)
            time_label = signup_val
            if precision == 'month' and not is_multiline:
                time_label = signup_val + '（月份）'
            obj = {
                'name': str(get_field(arr, 'name')).strip(),
                'level': str(get_field(arr, 'level')).strip(),
                'signupPrecision': precision,
                'signupTime': signup_val,
                'signupTimeLabel': time_label,
                'category': str(get_field(arr, 'category')).strip(),
                'organizer': str(get_field(arr, 'organizer')).strip(),
                'undertaker': str(get_field(arr, 'undertaker')).strip(),
                'signupLink': str(get_field(arr, 'signupLink')).strip(),
            }
            obj['tags'] = auto_suggest_tags_for_row(obj)
            if obj['name']:
                nkey = normalize_name(obj['name'])
                obj['exists'] = nkey in db_names
                obj['duplicate'] = nkey in file_seen
                file_seen.add(nkey)
                parsed.append(obj)

        return jsonify({
            'colMap': col_map,
            'headers': headers,
            'headerIdx': header_idx,
            'rows': parsed,
        })
    except Exception as e:
        return jsonify({'error': f'解析失败：{e}'}), 500


@app.route('/api/excel/template', methods=['GET'])
def excel_template():
    """下载模板 xls（使用项目内置的竞赛调研模板）"""
    template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '竞赛调研模板.xls')
    if not os.path.exists(template_path):
        return jsonify({'error': '模板文件不存在: ' + template_path}), 404
    return send_file(template_path, as_attachment=True, download_name='竞赛调研模板.xls',
                     mimetype='application/vnd.ms-excel')


if __name__ == '__main__':
    host = os.environ.get('HOST', '127.0.0.1')
    port = int(os.environ.get('PORT', '8765'))
    init_db()
    print(f'[server] DB: {DB_PATH}')
    print(f'[server] HTML: {HTML_FILE}')
    print(f'[server] Starting on http://{host}:{port}')
    app.run(host=host, port=port, debug=False)
